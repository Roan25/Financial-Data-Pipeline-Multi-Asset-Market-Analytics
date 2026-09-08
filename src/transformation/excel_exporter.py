import openpyxl
from openpyxl.styles import Font, PatternFill
import pandas as pd
import json

class ExcelReportGenerator:
    def __init__(self, db_manager):
        self.db = db_manager
        self.filepath = "data/processed/Nifty_Executive_Report.xlsx"
        import os
        os.makedirs(os.path.dirname(self.filepath), exist_ok=True)

    def generate(self, rejected_logs):
        wb = openpyxl.Workbook()
        
        # Sheet 1: Sector Rotation
        ws1 = wb.active
        ws1.title = "Sector Rotation"
        sectors_df = self.db.query("SELECT sector, SUM(market_cap) as tmc, AVG(pct_change) as ret FROM equity_snapshot GROUP BY sector")
        ws1.append(["Sector", "Total Market Cap", "Avg Return (%)"])
        for _, row in sectors_df.iterrows():
            ws1.append([row['sector'], row['tmc'], row['ret']])
        for cell in ws1[1]: cell.font = Font(bold=True)

        # Sheet 2: F&O Microstructure
        ws2 = wb.create_sheet(title="F&O Microstructure")
        fo_df = self.db.query("SELECT strike_price, option_type, ltp, open_interest, change_in_oi, market_phase FROM fo_contracts")
        ws2.append(list(fo_df.columns))
        
        fills = {
            "Long Buildup": PatternFill(start_color="C6EFCE", fill_type="solid"),
            "Short Buildup": PatternFill(start_color="FFC7CE", fill_type="solid"),
            "Long Unwinding": PatternFill(start_color="FFEB9C", fill_type="solid"),
            "Short Covering": PatternFill(start_color="B4C6E7", fill_type="solid")
        }
        
        for r_idx, row in enumerate(fo_df.itertuples(index=False), 2):
            ws2.append(list(row))
            phase = row[5]
            if phase in fills:
                for c_idx in range(1, 7):
                    ws2.cell(row=r_idx, column=c_idx).fill = fills[phase]

        # Sheet 3: Pipeline Audit Log
        ws3 = wb.create_sheet(title="Pipeline Audit Log")
        ws3.append(["Timestamp", "Reason", "Payload"])
        for log in rejected_logs:
            ws3.append([log['timestamp'], log['reason'], json.dumps(log['payload'], default=str)])

        wb.save(self.filepath)
